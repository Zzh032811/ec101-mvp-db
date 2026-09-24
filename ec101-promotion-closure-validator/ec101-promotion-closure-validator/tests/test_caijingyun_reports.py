import json
import unittest
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "outputs" / "勤进"


class CaijingyunReportTest(unittest.TestCase):
    def test_one_plus_four_reports_exist(self):
        expected = {
            "EC101平台活动闭环能力矩阵.xlsx",
            "单平台数据闭环验证报告.xlsx",
            "业务规则与待确认事项.xlsx",
            "最小数据要求与Gap清单.xlsx",
            "标准字段映射表.xlsx",
        }
        self.assertEqual({path.name for path in OUTPUT_DIR.glob("*.xlsx") if not path.name.startswith(".~")}, expected)

    def test_validation_report_uses_finance_cloud_results(self):
        workbook = load_workbook(OUTPUT_DIR / "单平台数据闭环验证报告.xlsx", data_only=True)
        summary = workbook["验证摘要"]
        self.assertEqual(summary["I7"].value, "未验证")
        self.assertEqual(summary["G7"].value, "已覆盖开始日后销售数据")
        self.assertEqual(summary["G8"].value, "99/99销售明细匹配")
        self.assertEqual(summary["A8"].value, "满赠")
        self.assertEqual(summary["D8"].value, "未验证")
        self.assertEqual(summary["E8"].value, "97个候选抱枕")
        self.assertIn("活动配置", workbook.sheetnames)
        self.assertIn("满减数据覆盖", workbook.sheetnames)
        self.assertIn("满赠逐单", workbook.sheetnames)

    def test_field_mapping_workbook_matches_run_json(self):
        run_data = json.loads((OUTPUT_DIR / "run.json").read_text(encoding="utf-8"))
        workbook = load_workbook(OUTPUT_DIR / "标准字段映射表.xlsx", data_only=True)
        sheet = workbook["字段映射"]
        self.assertEqual(sheet.max_row, len(run_data["field_mappings"]) + 6)
        rows = [tuple(sheet.cell(row, column).value for column in range(1, 11)) for row in range(7, sheet.max_row + 1)]
        self.assertTrue(any(row[2] == "销售-差价金额" for row in rows))
        self.assertTrue(any(row[3] == "订单状态" and row[8] == "Gap" for row in rows))


if __name__ == "__main__":
    unittest.main()
