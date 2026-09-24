import json
import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "outputs" / "顺英"


class ReportContractTest(unittest.TestCase):
    def test_one_plus_four_reports_exist(self):
        expected = {
            "EC101平台活动闭环能力矩阵.xlsx",
            "单平台数据闭环验证报告.xlsx",
            "业务规则与待确认事项.xlsx",
            "最小数据要求与Gap清单.xlsx",
            "标准字段映射表.xlsx",
        }
        actual = {path.name for path in OUTPUT_DIR.glob("*.xlsx") if not path.name.startswith(".~")}
        self.assertEqual(actual, expected)

    def test_order_ids_are_formatted_as_text_in_detail_sheets(self):
        workbook = load_workbook(OUTPUT_DIR / "单平台数据闭环验证报告.xlsx", data_only=True)
        for sheet_name in ("满减逐单", "满赠逐单"):
            cell = workbook[sheet_name]["A5"]
            self.assertIsInstance(cell.value, str)
            self.assertEqual(cell.number_format, "@")

    def test_field_mapping_report_contains_complete_union(self):
        run_data = json.loads((OUTPUT_DIR / "run.json").read_text(encoding="utf-8"))
        workbook = load_workbook(OUTPUT_DIR / "标准字段映射表.xlsx", data_only=True)
        sheet = workbook["字段映射"]

        expected_headers = [
            "模块",
            "来源文件",
            "经销商原始字段",
            "EC101标准字段",
            "字段关系",
            "来源类型",
            "重要度",
            "证据等级",
            "状态",
            "说明",
        ]
        actual_headers = [sheet.cell(6, column).value for column in range(1, 11)]
        self.assertEqual(actual_headers, expected_headers)
        self.assertEqual(sheet.max_row, len(run_data["field_mappings"]) + 6)

        rows = [
            tuple(sheet.cell(row, column).value for column in range(1, 11))
            for row in range(7, sheet.max_row + 1)
        ]
        self.assertTrue(any(row[2] == "登录账号" and row[4] == "非V1字段" for row in rows))
        self.assertTrue(any(row[3] == "活动编号" and row[8] == "Gap" for row in rows))


if __name__ == "__main__":
    unittest.main()
